"""The authority-ready dossier, as DOCX.

Sections, in order: cover, ConOps, operational volume, the SORA assessment tables
(iGRC → GRC → ARC → SAIL), the OSO compliance matrix, the mitigation evidence index, an
ERP skeleton, and the **justification appendix** — the full trace, every node with its
rule and citation.

Three rules this module keeps:

1. **Every risk value comes from the stored determination, verbatim.** Nothing here
   recomputes, rounds or re-derives. The tables are a rendering of the record.
2. **The disclaimer appears in every page header**, not only on the cover, because a
   printed page can be read on its own.
3. **An UNVERIFIED rulebase is stated on the cover and on every page.** A dossier built
   on seed tables must not be able to pass for one that was checked.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document as new_document
from docx.document import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services import mapimage

DISCLAIMER = (
    "SORA Copilot prepares and verifies your risk assessment. Your national aviation "
    "authority grants the authorisation."
)

UNVERIFIED_BANNER = (
    "UNVERIFIED RULEBASE — the regulatory tables behind this assessment have not been "
    "checked against the published JARUS SORA 2.5 text. This document must not be "
    "submitted to an authority."
)

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x60, 0x66)
WARN = RGBColor(0x8A, 0x52, 0x10)

ENVIRONMENT_LABELS = {
    "controlled_ground_area": "controlled ground area",
    "sparsely_populated": "sparsely populated",
    "populated": "populated",
    "assembly_of_people": "assembly of people",
}

MITIGATION_LABELS = {
    "m1_strategic": "M1 — sheltering",
    "m1_ers": "M1 — ground observation",
    "m2_impact": "M2 — impact reduction",
    "strategic_air": "Strategic air risk mitigation",
    "tactical_air": "Tactical air risk mitigation",
}


@dataclass(frozen=True, slots=True)
class DossierInput:
    """Everything the document renders. Assembled by the caller; nothing is fetched here."""

    operation: dict[str, Any]
    organisation: dict[str, Any]
    drone: dict[str, Any] | None
    mission: dict[str, Any]
    determination: dict[str, Any]
    evidence: list[dict[str, Any]]
    geo: dict[str, Any] | None
    version: int
    #: Optional LLM-written narrative. Absent when no key is configured, and the document
    #: states that rather than leaving a blank section.
    prose: dict[str, str] | None = None
    #: The derived outlines, for the static chart. Absent when nothing has been drawn.
    outlines: dict[str, Any] | None = None


def _style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    # East-Asian font name too, or Word substitutes its own for some glyphs.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def _header(document: Document, unverified: bool, stamp: str) -> None:
    """The disclaimer on every page, and the UNVERIFIED banner when it applies."""
    for section in document.sections:
        header = section.header
        header.is_linked_to_previous = False
        first = header.paragraphs[0]
        first.text = ""
        run = first.add_run(DISCLAIMER)
        run.font.size = Pt(7.5)
        run.font.color.rgb = MUTED
        first.paragraph_format.space_after = Pt(0)

        line = header.add_paragraph()
        run = line.add_run(stamp)
        run.font.size = Pt(7.5)
        run.font.color.rgb = MUTED

        if unverified:
            warning = header.add_paragraph()
            run = warning.add_run(UNVERIFIED_BANNER)
            run.font.size = Pt(7.5)
            run.bold = True
            run.font.color.rgb = WARN


def _h(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = INK
        run.font.name = "Calibri"


def _kv_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    table.columns[0].width = Cm(6.5)
    table.columns[1].width = Cm(10.5)
    document.add_paragraph()


def _note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def _sail_text(determination: dict[str, Any]) -> str:
    sail = determination.get("sail")
    return "Certified category" if sail == "certified_category" else f"SAIL {sail}"


def build(data: DossierInput) -> bytes:
    """Render the dossier. Returns the .docx bytes."""
    d = data.determination
    unverified = bool(d.get("warnings"))
    stamp = (
        f"Rulebase {d.get('rulebase_version')} · determination {str(d.get('id'))[:8]}"
        f" · dossier v{data.version}"
    )

    document = new_document()
    _style(document)

    section = document.sections[0]
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # ---------------------------------------------------------------- cover
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Specific Operations Risk Assessment")
    run.bold = True
    run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    run = subtitle.add_run(str(data.operation.get("name", "Operation")))
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED

    document.add_paragraph()
    _kv_table(
        document,
        [
            ("Operator", str(data.organisation.get("name", "—"))),
            ("Operator registration", str(data.organisation.get("operator_registration") or "—")),
            ("Competent authority", _authority(data.organisation)),
            ("Operation", str(data.operation.get("name", "—"))),
            ("Site", str(data.operation.get("site") or "—")),
            ("Outcome", _sail_text(d)),
            ("Rulebase version", str(d.get("rulebase_version"))),
            ("Computed at", str(d.get("computed_at"))),
            ("Mission hash", str(d.get("mission_hash", ""))[:32]),
            ("Dossier version", str(data.version)),
            ("Generated", datetime.now(UTC).strftime("%d %b %Y %H:%M UTC")),
        ],
    )

    if unverified:
        _h(document, "Verification status", 2)
        paragraph = document.add_paragraph()
        run = paragraph.add_run(UNVERIFIED_BANNER)
        run.bold = True
        run.font.color.rgb = WARN
        for warning in d.get("warnings", []):
            document.add_paragraph(str(warning), style="List Bullet")

    document.add_paragraph(DISCLAIMER)

    # ---------------------------------------------------------------- ConOps
    document.add_section(WD_SECTION.NEW_PAGE)
    _h(document, "1. Concept of operations")
    if data.prose and data.prose.get("conops"):
        for block in data.prose["conops"].split("\n\n"):
            if block.strip():
                document.add_paragraph(block.strip())
        _note(
            document,
            "Narrative generated from the facts and values above. No risk value in it "
            "was produced by a language model.",
        )
    else:
        _conops_from_facts(document, data)

    _h(document, "1.1 Aircraft", 2)
    drone = data.drone or {}
    _kv_table(
        document,
        [
            ("Model", str(drone.get("name", "—"))),
            ("Registration", str(drone.get("registration") or "—")),
            ("Maximum characteristic dimension", f"{drone.get('max_dimension_m', '—')} m"),
            ("MTOM", f"{drone.get('mtom_kg', '—')} kg"),
            ("Cruise speed", f"{drone.get('cruise_speed_ms', '—')} m/s"),
            ("Maximum speed", f"{drone.get('max_speed_ms', '—')} m/s"),
            (
                "Kinetic energy at maximum speed",
                f"{float(drone.get('kinetic_energy_j') or 0) / 1000:.1f} kJ",
            ),
            (
                "Fitted mitigations",
                ", ".join(drone.get("fitted_mitigations") or []) or "none",
            ),
        ],
    )

    # ---------------------------------------------------------------- volume
    _h(document, "2. Operational volume")
    volume = data.mission.get("volume") or {}
    rows = [
        ("Maximum altitude", f"{volume.get('max_altitude_agl_m', '—')} m AGL"),
        (
            "Altitude band",
            "above 500 ft (152 m) AGL"
            if float(volume.get("max_altitude_agl_m") or 0) > 152.4
            else "at or below 500 ft (152 m) AGL",
        ),
        ("Operation type", str(data.mission.get("operation_type", "—")).upper()),
    ]
    if data.geo:
        rows += [
            ("Flight geography", f"{data.geo.get('flight_geography_km2', '—')} km²"),
            ("Contingency volume", f"{data.geo.get('contingency_volume_km2', '—')} km²"),
            ("Ground risk buffer", f"{data.geo.get('ground_risk_buffer_km2', '—')} km²"),
            ("Adjacent area", f"{data.geo.get('adjacent_area_km2', '—')} km²"),
        ]
    _kv_table(document, rows)

    _volume_chart(document, data)

    if data.geo and data.geo.get("derivations"):
        _h(document, "2.2 How the volumes were derived", 2)
        for derivation in data.geo["derivations"]:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{derivation.get('value_m')} m — ").bold = True
            paragraph.add_run(str(derivation.get("rule", "")))
        _note(
            document,
            "Derived from the drawn flight geography in an equal-area projection. No "
            "volume in this dossier was entered by hand.",
        )
    else:
        _note(document, "The operational volumes have not been derived for this operation.")

    # ---------------------------------------------------------------- the chain
    _h(document, "3. SORA assessment")
    _kv_table(
        document,
        [
            ("Intrinsic ground risk class (iGRC)", _or_dash(d.get("igrc"))),
            ("Final ground risk class (GRC)", _or_dash(d.get("final_grc"))),
            ("Initial air risk class (ARC)", str(d.get("initial_arc", "—")).upper()),
            ("Residual air risk class", str(d.get("residual_arc", "—")).upper()),
            ("SAIL", _sail_text(d)),
            ("TMPR", str(d.get("tmpr", "—"))),
            ("Containment", _containment_text(d)),
        ],
    )

    _h(document, "3.1 Ground environment", 2)
    _kv_table(
        document,
        [
            ("Ground environment", _environment(data.mission)),
            (
                "Peak population density",
                f"{data.mission.get('population_density_per_km2', '—')} ppl/km²",
            ),
        ],
    )

    _h(document, "3.2 Mitigations claimed", 2)
    claims = data.mission.get("mitigations") or {}
    claimed = [(MITIGATION_LABELS.get(k, k), str(v)) for k, v in claims.items() if v and v != "none"]
    if claimed:
        _kv_table(document, claimed)
    else:
        document.add_paragraph("No mitigations claimed.")

    _h(document, "3.3 Airspace", 2)
    _kv_table(
        document,
        [
            ("Airspace class", str(data.mission.get("airspace_class", "—")).upper()),
            ("Controlled airspace", _yes_no(data.mission.get("controlled_airspace"))),
            ("Inside the aerodrome environment", _yes_no(data.mission.get("near_aerodrome"))),
            ("Atypical or segregated airspace", _yes_no(data.mission.get("atypical_airspace"))),
        ],
    )

    # ---------------------------------------------------------------- OSO matrix
    document.add_section(WD_SECTION.NEW_PAGE)
    _h(document, "4. Operational safety objectives")
    osos = d.get("osos") or []
    required = [o for o in osos if o.get("required_robustness") != "none"]
    document.add_paragraph(
        f"{len(required)} of {len(osos)} objectives apply at {_sail_text(d)}. "
        "Objectives that do not apply are listed so that it is clear they were considered."
    )

    by_oso: dict[str, list[dict[str, Any]]] = {}
    for item in data.evidence:
        by_oso.setdefault(str(item.get("oso_id") or ""), []).append(item)

    table = document.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(
        table.rows[0].cells, ["OSO", "Objective", "Required", "Status", "Evidence"], strict=False
    ):
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for oso in osos:
        cells = table.add_row().cells
        cells[0].text = str(oso.get("oso_id", ""))
        cells[1].text = str(oso.get("title", ""))
        level = str(oso.get("required_robustness", "none"))
        cells[2].text = "not applicable" if level == "none" else level
        cells[3].text = "not applicable" if level == "none" else str(oso.get("status", "unmet"))
        files = by_oso.get(str(oso.get("oso_id", "")), [])
        cells[4].text = "\n".join(str(f.get("filename")) for f in files) if files else "—"

    document.add_paragraph()
    _note(
        document,
        "Required robustness is assigned by the rulebase from the computed SAIL. Evidence "
        "status reflects what has been attached in SORA Copilot, not an authority's view "
        "of its sufficiency.",
    )

    # ---------------------------------------------------------------- evidence index
    _h(document, "5. Mitigation and evidence index")
    if data.evidence:
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(
            table.rows[0].cells, ["File", "Attached to", "Size", "Uploaded"], strict=False
        ):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for item in data.evidence:
            cells = table.add_row().cells
            cells[0].text = str(item.get("filename", ""))
            cells[1].text = str(item.get("oso_id") or "the operation")
            size = item.get("bytes")
            cells[2].text = f"{int(size) / 1024:.0f} kB" if size else "—"
            cells[3].text = str(item.get("uploaded_at", ""))[:10]
        document.add_paragraph()
    else:
        document.add_paragraph("No evidence files have been attached to this operation.")

    # ---------------------------------------------------------------- ERP
    _h(document, "6. Emergency response plan")
    if data.prose and data.prose.get("erp"):
        for block in data.prose["erp"].split("\n\n"):
            if block.strip():
                document.add_paragraph(block.strip())
    else:
        document.add_paragraph(
            "This section is a skeleton. SORA Copilot does not write an emergency response "
            "plan for you: it depends on the site, the crew and the local services, none of "
            "which the tool knows. The headings below are the ones an authority will expect."
        )
        for heading in (
            "Alerting — who is called, in what order, with what information",
            "Containment of the aircraft after a loss of control",
            "Casualty response and first aid provision on site",
            "Securing the area and preserving evidence",
            "Notification of the competent authority and occurrence reporting",
            "Crew debrief and return to service",
        ):
            document.add_paragraph(heading, style="List Bullet")

    # ---------------------------------------------------------------- the appendix
    document.add_section(WD_SECTION.NEW_PAGE)
    _h(document, "Appendix A — Justification trace")
    document.add_paragraph(
        "Every conclusion the rules engine reached, the rule that produced it, and the "
        "clause that rule comes from. The engine cannot produce a conclusion without a "
        "justification: a model in which one is missing is rejected as unsatisfiable."
    )

    trace = d.get("trace") or []
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(
        table.rows[0].cells, ["Conclusion", "Rule", "Citation", "From"], strict=False
    ):
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for node in trace:
        cells = table.add_row().cells
        cells[0].text = str(node.get("conclusion", ""))
        cells[1].text = str(node.get("rule_id", ""))
        citation = node.get("citation") or {}
        suffix = "" if citation.get("verified") else "  [UNVERIFIED]"
        page = f", p. {citation.get('page')}" if citation.get("page") else ""
        cells[2].text = f"{citation.get('document', '')} · {citation.get('section', '')}{page}{suffix}"
        cells[3].text = ", ".join(node.get("premises") or []) or "—"

    document.add_paragraph()
    _note(document, f"{len(trace)} nodes. Mission hash {d.get('mission_hash', '')}.")

    _h(document, "Appendix B — Rulebase provenance", 1)
    document.add_paragraph(
        f"This determination was computed against rulebase {d.get('rulebase_version')}. "
        "A determination is pinned to the rulebase version it was computed against, so a "
        "later change to the rulebase does not alter this record."
    )
    if unverified:
        for warning in d.get("warnings", []):
            document.add_paragraph(str(warning), style="List Bullet")

    _header(document, unverified, stamp)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _volume_chart(document: Document, data: DossierInput) -> None:
    """The volumes as a chart, drawn from the same geometry the engine read.

    Not a basemap: a dossier is printed and archived, and a third-party tile would be a
    dead reference in a year. What matters in the document is the relationship between
    the volumes and their scale, which this shows exactly.
    """
    outlines = data.outlines or {}
    flight_geography = (data.mission.get("volume") or {}).get("flight_geography")
    if not flight_geography:
        return
    png = mapimage.render(
        flight_geography=flight_geography,
        contingency_volume=outlines.get("contingency_volume"),
        ground_risk_buffer=outlines.get("ground_risk_buffer"),
        adjacent_area=outlines.get("adjacent_area"),
    )
    if png is None:
        return
    _h(document, "2.1 Operational volumes", 2)
    document.add_picture(BytesIO(png), width=Cm(16.0))
    _note(
        document,
        "Drawn from the submitted flight geography and the volumes derived from it. "
        "North is up; the scale bar matches the grid spacing.",
    )


def _conops_from_facts(document: Document, data: DossierInput) -> None:
    """The ConOps written from the facts, with no language model involved."""
    mission = data.mission
    drone = data.drone or {}
    volume = mission.get("volume") or {}
    environment = ENVIRONMENT_LABELS.get(
        str(mission.get("ground_environment")), str(mission.get("ground_environment", "—"))
    )
    document.add_paragraph(
        f"{data.organisation.get('name', 'The operator')} will conduct "
        f"{str(mission.get('operation_type', '—')).upper()} operations with a "
        f"{drone.get('name', 'the aircraft')}"
        f"{f' ({drone.get('registration')})' if drone.get('registration') else ''} "
        f"at or below {volume.get('max_altitude_agl_m', '—')} m AGL over a {environment} "
        f"ground environment."
    )
    if mission.get("population_density_per_km2") is not None:
        document.add_paragraph(
            f"The peak population density recorded across the ground risk buffer is "
            f"{mission.get('population_density_per_km2')} people per square kilometre, "
            "which is the row used for the intrinsic ground risk class."
        )
    document.add_paragraph(
        f"The operation is planned in class {str(mission.get('airspace_class', '—')).upper()} "
        f"airspace. It is {'inside' if mission.get('near_aerodrome') else 'outside'} an "
        f"aerodrome environment and "
        f"{'in' if mission.get('controlled_airspace') else 'outside'} controlled airspace."
    )
    _note(
        document,
        "Written from the confirmed facts. No language model contributed to this section, "
        "and no language model contributed any risk value anywhere in this document.",
    )


def _environment(mission: dict[str, Any]) -> str:
    raw = str(mission.get("ground_environment", "—"))
    return ENVIRONMENT_LABELS.get(raw, raw)


def _or_dash(value: Any) -> str:
    return "—" if value is None else str(value)


def _yes_no(value: Any) -> str:
    return "yes" if value else "no"


def _containment_text(determination: dict[str, Any]) -> str:
    containment = determination.get("containment") or {}
    if containment.get("assessment_required"):
        return "assessment required"
    return str(containment.get("level") or "none")


def _authority(organisation: dict[str, Any]) -> str:
    return {
        "AT": "Austro Control (AT)",
        "DE": "Luftfahrt-Bundesamt (DE)",
        "EU": "the competent authority of the state of registration",
    }.get(str(organisation.get("country", "AT")), "—")
